"""Shared helpers for platform parsers."""
from openpyxl import load_workbook


def load_ws(filepath, sheet_name):
    """Load one sheet's rows as a list of tuples (read-only, values only)."""
    wb = load_workbook(filepath, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        return None
    ws = wb[sheet_name]
    return list(ws.iter_rows(values_only=True))


def sheet_names(filepath):
    wb = load_workbook(filepath, read_only=True)
    return wb.sheetnames


def to_num(v, default=0):
    if v is None or v == "":
        return default
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


class Recruiter:
    """One recruiter's activity row extracted from a raw export."""
    __slots__ = ("name", "searches", "profiles_viewed", "outreach_sent", "responses_received", "notes")

    def __init__(self, name, searches=0, profiles_viewed=0, outreach_sent=0, responses_received=0, notes=""):
        self.name = name
        self.searches = searches
        self.profiles_viewed = profiles_viewed
        self.outreach_sent = outreach_sent
        self.responses_received = responses_received
        self.notes = notes

    def as_dict(self):
        return {
            "name": self.name, "searches": self.searches, "profiles_viewed": self.profiles_viewed,
            "outreach_sent": self.outreach_sent, "responses_received": self.responses_received,
            "notes": self.notes,
        }


class ParseResult:
    """Standard output shape every platform parser returns."""

    def __init__(self, platform, recruiters=None, platform_totals=None, cumulative_fields=None, source_note=""):
        self.platform = platform
        self.recruiters = recruiters or []          # list[Recruiter]
        self.platform_totals = platform_totals or {}  # e.g. {"applications": 1200, "contacts_made": 300}
        # names of keys in platform_totals (and per-recruiter fields) that are CUMULATIVE-since-start
        # rather than already scoped to this reporting period. The pipeline diffs these against history.
        self.cumulative_fields = cumulative_fields or set()
        self.source_note = source_note
